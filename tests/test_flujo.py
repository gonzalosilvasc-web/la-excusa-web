"""Pruebas funcionales del sistema (se ejecutan con: python tests/test_flujo.py).

Cubren el flujo completo: estructura, mandatos, plantillas, casos,
clasificación, generación, revisor cuerpo–petitorio, bloqueos, auditoría,
seguridad de escritura y modo sin API key.
"""

from __future__ import annotations

import os
import shutil
import sys
from datetime import date
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

# Asegurar modo sin API key para las pruebas
for var in ("OPENAI_API_KEY", "GOOGLE_API_KEY", "OCR_PROVIDER"):
    os.environ.pop(var, None)

import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402

from legal import casos, clasificador, extraccion, generador, mandatos, plantillas, revisor  # noqa: E402
from legal.config import (  # noqa: E402
    BACKUPS_DIR, DATABASE_DIR, GENERATED_DIR, MANDATOS_XLSX, TEMP_DIR,
)
from legal.storage import (  # noqa: E402
    ExcelAbiertoError, ensure_structure, leer_auditoria, registrar_auditoria,
    write_excel_safe,
)
from legal.textutil import normalizar_nombre, parsear_monto  # noqa: E402

RESULTADOS: list[tuple[str, bool, str]] = []


def check(nombre: str, condicion: bool, detalle: str = "") -> None:
    RESULTADOS.append((nombre, condicion, detalle))
    print(f"{'  [PASA] ' if condicion else '  [FALLA]'} {nombre}"
          + (f" — {detalle}" if detalle and not condicion else ""))
    if not condicion:
        raise AssertionError(f"{nombre}: {detalle}")


def preparar_base_limpia() -> None:
    if DATABASE_DIR.exists():
        shutil.rmtree(DATABASE_DIR)
    ensure_structure()


def crear_pdf_falso(nombre: str) -> Path:
    TEMP_DIR.mkdir(parents=True, exist_ok=True)
    ruta = TEMP_DIR / nombre
    ruta.write_bytes(b"%PDF-1.4 test\n%%EOF")
    return ruta


def crear_docx_con_placeholders(nombre: str) -> Path:
    return plantillas.crear_plantilla_base(nombre)


def crear_docx_sin_placeholders(nombre: str) -> Path:
    ruta = TEMP_DIR / nombre
    doc = Document()
    doc.add_paragraph("Modelo antiguo del estudio sin variables.")
    doc.save(str(ruta))
    return ruta


# ---------------------------------------------------------------------------
print("\n== 1. Estructura y utilidades ==")
preparar_base_limpia()
check("Estructura creada automáticamente",
      MANDATOS_XLSX.exists() and (DATABASE_DIR / "cases.xlsx").exists()
      and (DATABASE_DIR / "pdf_storage").is_dir())
check("Normalización de nombres",
      normalizar_nombre("  MARÍA  Fernanda   SOTO-Pérez ") == "maria fernanda soto perez")
check("Parseo de montos chilenos",
      parsear_monto("$ 12.345.678") == 12345678 and parsear_monto("1.234,56") == 1234.56)

# ---------------------------------------------------------------------------
print("\n== 2. Modo sin API key ==")
disp, prov = extraccion.ocr_disponible()
check("OCR reporta no disponible sin API key", disp is False)
res = extraccion.extract_document_data("no_existe.pdf", "pagare")
check("Extracción sin API devuelve mensaje de modo manual",
      "_error" in res and "manualmente" in res["_error"])

# ---------------------------------------------------------------------------
print("\n== 3. Mandatos ==")
pdf = crear_pdf_falso("mandato_prueba.pdf")
reg_m = mandatos.guardar_mandato({
    "Nombre_Ejecutivo": "María Fernanda Soto Pérez",
    "Entidad_Representada": "Banco Santander",
    "Fecha_Otorgamiento": "15/03/2019",
    "Fecha_Revocacion": None,
    "Notario_Titular": "Juan San Martín",
    "Ciudad_Notaria": "Santiago",
    "Repertorio": "12345-2019",
    "Facultad_Pagare": True,
    "Texto_Facultad_Pagare": "Podrá suscribir pagarés…",
    "Usuario_Carga": "tester",
}, pdf)
mandatos.guardar_mandato({
    "Nombre_Ejecutivo": "Pedro Rojas Díaz",
    "Entidad_Representada": "Banco Santander",
    "Fecha_Otorgamiento": "10/01/2018",
    "Fecha_Revocacion": "20/06/2020",
    "Notario_Titular": "Iván Torrealba",
    "Ciudad_Notaria": "Santiago",
    "Repertorio": "999-2018",
    "Facultad_Pagare": True,
    "Usuario_Carga": "tester",
})
df_m = mandatos.cargar_mandatos()
check("Mandato guardado en mandatos.xlsx", len(df_m) == 2)
check("PDF del mandato guardado",
      (BASE / str(reg_m["PDF_Path"])).exists())

r = mandatos.buscar_personeria(df_m, "soto perez maria", "Banco Santander", date(2020, 5, 1))
check("Búsqueda fuzzy encuentra mandato vigente", r["vigente"] is not None)
texto = mandatos.texto_personeria(r["vigente"])
check("Texto de personería generado",
      "María Fernanda Soto Pérez" in texto and "12345-2019" in texto
      and "repertorio" in texto)
r2 = mandatos.buscar_personeria(df_m, "Pedro Rojas Diaz", "Banco Santander", date(2021, 1, 1))
check("Mandato revocado no aparece vigente tras revocación", r2["vigente"] is None)
r3 = mandatos.buscar_personeria(df_m, "Pedro Rojas Diaz", "Banco Santander", date(2020, 6, 20))
check("Advertencia de revocación el mismo día del instrumento",
      not r3["mismo_dia"].empty and r3["vigente"] is None)
dups = mandatos.detectar_duplicados(
    df_m, "Banco Santander", "maria fernanda soto perez",
    pd.Timestamp("2019-03-15"), "12345-2019")
check("Detección de duplicados de mandato", len(dups) == 1)

# ---------------------------------------------------------------------------
print("\n== 4. Plantillas Word ==")
tpl_reg = plantillas.registrar_plantilla_base_si_falta()
check("Plantilla base de fábrica registrada", tpl_reg is not None)
df_t = plantillas.cargar_templates()
phs = str(df_t.iloc[0]["Placeholders_Detectados"])
check("Placeholders detectados en plantilla con variables",
      "PETITORIO" in phs and "TITULOS_EJECUTIVOS" in phs and "TEXTO_PERSONERIA" in phs)

sin_ph = crear_docx_sin_placeholders("modelo_sin_placeholders.docx")
reg_sin = plantillas.guardar_template(sin_ph, {
    "Nombre_Template": "Modelo antiguo sin variables",
    "Tipo_Demanda": "Otros", "Usuario_Carga": "tester", "Activo": True,
})
check("Modelo sin placeholders queda BLOQUEADO para generación automática",
      str(reg_sin["Activo"]).lower() in ("false", "0")
      and reg_sin["Placeholders_Detectados"] == "")
copia = plantillas.crear_copia_con_placeholders(BASE / str(reg_sin["Ruta_Template"]))
check("Se crea copia con placeholders sugeridos",
      copia.exists() and len(plantillas.detectar_placeholders(copia)) > 10)

# ---------------------------------------------------------------------------
print("\n== 5. Caso simple: pagaré en pesos ==")
caso = casos.crear_caso({
    "Banco_Entidad": "Banco Santander",
    "Cliente_Deudor": "Comercial Los Álamos Ltda",
    "RUT_Deudor": "76.111.222-3", "Usuario_Carga": "tester",
})
id_caso = caso["ID_Caso"]
check("Caso creado con carpeta propia", (BASE / caso["Ruta_Carpeta_Caso"]).is_dir())
doc_guardado = casos.guardar_documento_caso(
    id_caso, "pagaré banco.pdf", b"%PDF-1.4", "Pagaré")
check("Documento del caso guardado en su carpeta", doc_guardado.exists())

inst1 = casos.agregar_instrumento(id_caso, {
    "Tipo_Instrumento": "Pagare", "Numero_Instrumento": "296134",
    "Fecha_Suscripcion": pd.Timestamp("2020-05-10"),
    "Fecha_Mora": pd.Timestamp("2024-02-01"),
    "Banco_Acreedor": "Banco Santander",
    "Deudor_Principal": "Comercial Los Álamos Ltda",
    "RUT_Deudor_Principal": "76.111.222-3",
    "Monto_Original": "20000000", "Saldo_Insoluto": "15000000",
    "Moneda": "CLP", "Tasa_Interes": "1,2% mensual", "Interes_Penal": "máximo convencional",
    "Clausula_Aceleracion": True, "Liberacion_Protesto": True,
    "Firmas_Autorizadas_Notarialmente": True,
    "Fuente_Documento": "pagare_296134.pdf", "Fuente_Pagina": "1",
    "Fuente_Texto": "…prometo pagar…",
})
casos.agregar_parte(id_caso, {
    "Nombre": "Comercial Los Álamos Ltda", "RUT": "76.111.222-3",
    "Tipo_Parte": "DEMANDADO", "Calidad_Juridica": "Deudor principal",
    "Responde_Por_Instrumento": "TODOS", "Monto_Responsabilidad": "15000000",
    "Moneda_Responsabilidad": "CLP", "Domicilio": "Av. Providencia 1234, Santiago",
    "Fuente_Documento": "pagare_296134.pdf",
})
df_i = casos.cargar_instrumentos(id_caso)
df_p = casos.cargar_partes(id_caso)
check("Instrumento y parte registrados", len(df_i) == 1 and len(df_p) == 1)

clas = clasificador.clasificar_caso(df_i, df_p)
check("Clasificador sugiere 'Cobro de pagaré en pesos'",
      clas["tipo_sugerido"] == "Cobro de pagaré en pesos"
      and clas["confianza"] == "alta")

sug = clasificador.sugerir_template(
    plantillas.cargar_templates(), clas["tipo_sugerido"], "Banco Santander", "CLP")
check("Sugerencia de modelo compatible", sug["sugerido"] is not None)

mandato_dict = mandatos.cargar_mandatos().iloc[0].to_dict()


def texto_render(ctx: dict) -> str:
    """Renderiza el modelo base con el contexto y devuelve el texto del Word,
    replicando lo que hace la app antes de revisar."""
    ruta = BASE / str(plantillas.cargar_templates().iloc[0]["Ruta_Template"])
    salida = generador.generar_demanda_docx(ruta, ctx, "_render_revision.docx")
    return plantillas.texto_completo_docx(salida)


campos = {
    "TRIBUNAL": "S.J.L. en lo Civil de Santiago",
    "MATERIA": "Cobro de pagaré en pesos",
    "DEMANDANTE": "Banco Santander", "RUT_DEMANDANTE": "97.036.000-K",
    "REPRESENTANTE_DEMANDANTE": "María Fernanda Soto Pérez",
    "BIENES_EMBARGO": "Bienes muebles del domicilio de la demandada",
    "EQUIVALENCIA_PESOS": "", "GARANTIAS": "", "DOCUMENTOS_EXTRA": "",
}
ctx = generador.componer_contexto(caso, df_i, df_p, mandato_dict, campos)
check("Contexto compuesto con petitorio y títulos",
      "296134" in ctx["TITULOS_EJECUTIVOS"] and "296134" in ctx["DOCUMENTOS_PRIMER_OTROSI"]
      and "Comercial Los" in ctx["PETITORIO"] and "costas" in ctx["PETITORIO"])
check("Suma de capital correcta", "15.000.000" in ctx["MONTO_CAPITAL"])

ruta_tpl = BASE / str(df_t.iloc[0]["Ruta_Template"])
preview = generador.generar_demanda_docx(ruta_tpl, ctx, "preview_test.docx")
texto_final = plantillas.texto_completo_docx(preview)
rev = revisor.revisar_demanda(ctx, df_i, df_p, mandato_dict, texto_final)
check("Caso simple correcto: revisión SIN errores críticos",
      not rev.bloquea_descarga,
      "; ".join(h.explicacion for h in rev.criticos))
check("No quedan placeholders en el Word generado", "{{" not in texto_final)

nombre_final = generador.nombre_archivo_demanda(
    "Banco Santander", "Comercial Los Álamos Ltda",
    "Cobro de pagaré en pesos", id_caso)
archivo_final = generador.generar_demanda_docx(ruta_tpl, ctx, nombre_final)
reg_dda = generador.registrar_demanda_generada(
    id_caso, str(df_t.iloc[0]["ID_Template"]), str(mandato_dict["ID_Mandato"]),
    "Cobro de pagaré en pesos", archivo_final, ctx, "tester", 0,
    len(rev.advertencias), True)
check("Demanda .docx generada y guardada en generated_demands/",
      archivo_final.exists() and archivo_final.parent == GENERATED_DIR)
check("Demanda registrada con contexto JSON",
      (BASE / str(reg_dda["Contexto_JSON"])).exists())
registrar_auditoria("GENERACION_DEMANDA", "tester", id_caso=id_caso,
                    resultado="OK", observaciones=nombre_final)

# ---------------------------------------------------------------------------
print("\n== 6. Caso con múltiples pagarés y avalista limitado ==")
caso2 = casos.crear_caso({
    "Banco_Entidad": "Banco Santander",
    "Cliente_Deudor": "Transportes del Sur SpA", "Usuario_Carga": "tester",
})
id2 = caso2["ID_Caso"]
i1 = casos.agregar_instrumento(id2, {
    "Tipo_Instrumento": "Pagare", "Numero_Instrumento": "111",
    "Fecha_Suscripcion": pd.Timestamp("2021-03-01"),
    "Fecha_Mora": pd.Timestamp("2024-01-15"),
    "Banco_Acreedor": "Banco Santander",
    "Deudor_Principal": "Transportes del Sur SpA",
    "RUT_Deudor_Principal": "77.222.333-4",
    "Saldo_Insoluto": "10000000", "Moneda": "CLP",
    "Fuente_Documento": "p111.pdf",
})
i2 = casos.agregar_instrumento(id2, {
    "Tipo_Instrumento": "Pagare", "Numero_Instrumento": "222",
    "Fecha_Suscripcion": pd.Timestamp("2022-07-01"),
    "Fecha_Mora": pd.Timestamp("2024-01-15"),
    "Banco_Acreedor": "Banco Santander",
    "Deudor_Principal": "Transportes del Sur SpA",
    "RUT_Deudor_Principal": "77.222.333-4",
    "Saldo_Insoluto": "5000000", "Moneda": "CLP",
    "Fuente_Documento": "p222.pdf",
})
casos.agregar_parte(id2, {
    "Nombre": "Transportes del Sur SpA", "RUT": "77.222.333-4",
    "Calidad_Juridica": "Deudor principal",
    "Responde_Por_Instrumento": "TODOS", "Monto_Responsabilidad": "15000000",
    "Moneda_Responsabilidad": "CLP", "Domicilio": "Ruta 5 Sur km 10, Temuco",
    "Fuente_Documento": "p111.pdf",
})
casos.agregar_parte(id2, {
    "Nombre": "Rosa Sottorff Muñoz", "RUT": "9.999.999-9",
    "Calidad_Juridica": "Avalista",
    "Responde_Por_Instrumento": i1["ID_Instrumento"],
    "Monto_Responsabilidad": "10000000", "Moneda_Responsabilidad": "CLP",
    "Domicilio": "Los Aromos 45, Temuco", "Fuente_Documento": "p111.pdf",
})
df_i2 = casos.cargar_instrumentos(id2)
df_p2 = casos.cargar_partes(id2)
clas2 = clasificador.clasificar_caso(df_i2, df_p2)
check("Clasificador detecta múltiples pagarés",
      clas2["tipo_sugerido"] == "Cobro de múltiples pagarés")

ctx2 = generador.componer_contexto(caso2, df_i2, df_p2, mandato_dict, {
    "TRIBUNAL": "S.J.L. en lo Civil de Temuco",
    "MATERIA": "Cobro de múltiples pagarés",
    "DEMANDANTE": "Banco Santander", "RUT_DEMANDANTE": "97.036.000-K",
    "REPRESENTANTE_DEMANDANTE": "María Fernanda Soto Pérez",
    "BIENES_EMBARGO": "Camión patente XX-1234",
})
check("Suma de múltiples pagarés correcta", "15.000.000" in ctx2["MONTO_CAPITAL"])
check("Petitorio limita al avalista al pagaré que garantizó",
      "limitada exclusivamente" in ctx2["PETITORIO"] and "111" in ctx2["PETITORIO"])
rev2 = revisor.revisar_demanda(ctx2, df_i2, df_p2, mandato_dict, texto_render(ctx2))
check("Caso múltiple bien estructurado: sin errores críticos",
      not rev2.bloquea_descarga,
      "; ".join(h.explicacion for h in rev2.criticos))

# Avalista pedida por un pagaré que NO garantizó (monto superior a su tope)
df_p2_mal = df_p2.copy()
df_p2_mal["Monto_Responsabilidad"] = df_p2_mal["Monto_Responsabilidad"].astype(object)
df_p2_mal.loc[df_p2_mal["Nombre"] == "Rosa Sottorff Muñoz",
              "Monto_Responsabilidad"] = "15000000"
rev2b = revisor.revisar_demanda(ctx2, df_i2, df_p2_mal, mandato_dict, texto_render(ctx2))
check("Bloquea si al avalista se le pide más de lo que garantizó",
      any(h.regla == "responsabilidad" for h in rev2b.criticos))

# ---------------------------------------------------------------------------
print("\n== 7. Contradicción cuerpo–petitorio ==")
ctx3 = dict(ctx2)
# a) Rosa aparece en el cuerpo pero se elimina del petitorio
ctx3["PETITORIO"] = (
    "RUEGO A US. despachar mandamiento de ejecución y embargo en contra de "
    "Transportes del Sur SpA por $ 15.000.000, más intereses y costas."
)
rev3 = revisor.revisar_demanda(ctx3, df_i2, df_p2, mandato_dict, texto_render(ctx3))
err_rosa = [h for h in rev3.criticos if "Rosa Sottorff" in h.explicacion
            and h.regla == "cuerpo-petitorio"]
check("Bloquea: avalista en el cuerpo pero ausente del petitorio",
      bool(err_rosa) and rev3.bloquea_descarga)
print(f"    Mensaje: {err_rosa[0].explicacion}")

# b) Persona demandada en el petitorio sin respaldo en el cuerpo
ctx4 = dict(ctx2)
ctx4["DEMANDADOS"] = "1. Transportes del Sur SpA, RUT 77.222.333-4…"
ctx4["TITULOS_EJECUTIVOS"] = ctx4["TITULOS_EJECUTIVOS"].replace("Rosa Sottorff Muñoz", "")
ctx4["RELATO_DE_LOS_HECHOS"] = ctx4["RELATO_DE_LOS_HECHOS"].replace("Rosa Sottorff Muñoz", "")
rev4 = revisor.revisar_demanda(ctx4, df_i2, df_p2, mandato_dict, texto_render(ctx4))
check("Bloquea: demandada en el petitorio sin respaldo en el cuerpo",
      any("respaldo" in h.explicacion.lower() and "Rosa Sottorff" in h.explicacion
          for h in rev4.criticos))

# ---------------------------------------------------------------------------
print("\n== 8. Placeholders pendientes y comentarios internos ==")
ctx5 = dict(ctx2)
ctx5["RELATO_DE_LOS_HECHOS"] += "\n{{FALTA_ESTE_DATO}}"
rev5 = revisor.revisar_demanda(ctx5, df_i2, df_p2, mandato_dict, texto_render(ctx5))
check("Bloquea placeholder sin completar e indica cuál",
      any("FALTA_ESTE_DATO" in h.explicacion for h in rev5.criticos))

for marca in ("[EA1.1]", "[COMPLETAR]", "xxx", "pendiente", "revisar"):
    ctx6 = dict(ctx2)
    ctx6["RELATO_DE_LOS_HECHOS"] += f"\nnota {marca} interna"
    rev6 = revisor.revisar_demanda(ctx6, df_i2, df_p2, mandato_dict, texto_render(ctx6))
    check(f"Bloquea comentario interno «{marca}»",
          any(h.regla == "texto-prohibido" for h in rev6.criticos))

# Campo crítico vacío (BIENES_EMBARGO NO_DETECTADO)
ctx7 = generador.componer_contexto(caso2, df_i2, df_p2, mandato_dict, {
    "TRIBUNAL": "S.J.L. Temuco", "DEMANDANTE": "Banco Santander",
    "REPRESENTANTE_DEMANDANTE": "María Soto",
})
rev7 = revisor.revisar_demanda(ctx7, df_i2, df_p2, mandato_dict, "")
check("Bloquea campo crítico NO_DETECTADO (bienes embargo)",
      any(h.ubicacion == "BIENES_EMBARGO" for h in rev7.criticos))

# Personería posterior al título
mandato_tardio = dict(mandato_dict)
mandato_tardio["Fecha_Otorgamiento"] = pd.Timestamp("2023-01-01")
rev8 = revisor.revisar_demanda(ctx2, df_i2, df_p2, mandato_tardio, texto_render(ctx2))
check("Bloquea personería otorgada después del título",
      any(h.regla == "personeria-fecha" and h.nivel == revisor.CRITICO
          for h in rev8.hallazgos))

# Moneda USD sin equivalencia
df_i_usd = df_i2.copy()
df_i_usd["Moneda"] = "USD"
ctx9 = generador.componer_contexto(caso2, df_i_usd, df_p2, mandato_dict, {
    "TRIBUNAL": "S.J.L. Temuco", "DEMANDANTE": "Banco Santander",
    "REPRESENTANTE_DEMANDANTE": "María Soto", "BIENES_EMBARGO": "Camión",
})
rev9 = revisor.revisar_demanda(ctx9, df_i_usd, df_p2, mandato_dict, "")
check("Bloquea USD sin equivalencia en pesos",
      any(h.regla == "moneda" for h in rev9.criticos))

# ---------------------------------------------------------------------------
print("\n== 9. Excel abierto / bloqueado ==")
import legal.storage as storage_mod

_original_replace = os.replace


def _replace_bloqueado(src, dst, *a, **k):  # noqa: ANN001
    if str(dst).endswith("mandatos.xlsx"):
        raise PermissionError(13, "El proceso no puede obtener acceso al archivo")
    return _original_replace(src, dst, *a, **k)


os.replace = _replace_bloqueado
try:
    write_excel_safe(pd.DataFrame({"a": [1]}), MANDATOS_XLSX)
    check("Excel bloqueado lanza ExcelAbiertoError", False, "no lanzó excepción")
except ExcelAbiertoError as exc:
    check("Excel bloqueado lanza ExcelAbiertoError con mensaje claro",
          "Cierre el archivo" in str(exc))
finally:
    os.replace = _original_replace
df_verif = mandatos.cargar_mandatos()
check("El archivo original quedó intacto tras el bloqueo", len(df_verif) == 2)

# ---------------------------------------------------------------------------
print("\n== 10. Auditoría y respaldos ==")
log = leer_auditoria()
acciones = set(log["Accion"].astype(str))
check("Auditoría registra generación de demanda", "GENERACION_DEMANDA" in acciones)
registrar_auditoria("GENERACION_BLOQUEADA_REVISION", "tester", id_caso=id2,
                    resultado="BLOQUEADO", observaciones="prueba de bloqueo")
registrar_auditoria("DESCARGA_DEMANDA", "tester", id_caso=id_caso, resultado="OK")
log = leer_auditoria()
check("Auditoría registra bloqueos y descargas",
      "GENERACION_BLOQUEADA_REVISION" in set(log["Accion"])
      and "DESCARGA_DEMANDA" in set(log["Accion"]))
backups = list(BACKUPS_DIR.glob("*.xlsx"))
check("Respaldos automáticos creados con timestamp", len(backups) > 5)

# ---------------------------------------------------------------------------
print("\n== 11. Clasificador: variantes ==")
df_i_usd2 = df_i.copy()
df_i_usd2["Moneda"] = "USD"
c = clasificador.clasificar_caso(df_i_usd2, df_p)
check("Detecta pagaré en dólares", c["tipo_sugerido"] == "Cobro de pagaré en dólares")

df_i_mutuo = df_i.copy()
df_i_mutuo["Tipo_Instrumento"] = "Mutuo"
df_i_mutuo["Observaciones"] = "con garantía hipotecaria"
c = clasificador.clasificar_caso(df_i_mutuo, df_p)
check("Detecta mutuo hipotecario", c["tipo_sugerido"] == "Cobro de mutuo hipotecario")

df_p_tp = df_p.copy()
df_p_tp.loc[:, "Calidad_Juridica"] = "Tercero poseedor"
c = clasificador.clasificar_caso(df_i_mutuo, df_p_tp)
check("Detecta desposeimiento de tercer poseedor",
      c["tipo_sugerido"] == "Desposeimiento de tercer poseedor")

c = clasificador.clasificar_caso(pd.DataFrame(columns=df_i.columns), df_p)
check("Bloquea caso sin títulos ejecutivos", c["bloqueo"] is True)

# ---------------------------------------------------------------------------
print("\n== 12. Seguridad jurídica: nombre ambiguo entre dos personas ==")
mandatos.guardar_mandato({
    "Nombre_Ejecutivo": "María González Díaz",
    "Entidad_Representada": "Banco Santander",
    "Fecha_Otorgamiento": "01/06/2019", "Fecha_Revocacion": None,
    "Notario_Titular": "Notario Prueba", "Ciudad_Notaria": "Santiago",
    "Repertorio": "777-2019", "Facultad_Pagare": True, "Usuario_Carga": "tester",
})
df_amb = mandatos.cargar_mandatos()
r_amb = mandatos.buscar_personeria(df_amb, "maria", "Banco Santander", date(2020, 5, 1))
check("Nombre ambiguo (dos Marías): NO elige ninguna y reporta ambigüedad",
      r_amb["vigente"] is None and len(r_amb["ambiguos"]) >= 2)
r_ok = mandatos.buscar_personeria(
    df_amb, "María Fernanda Soto Pérez", "Banco Santander", date(2020, 5, 1))
check("Nombre completo resuelve la ambigüedad",
      r_ok["vigente"] is not None and r_ok["ambiguos"].empty
      and r_ok["vigente"]["Repertorio"] == "12345-2019")

# ---------------------------------------------------------------------------
print("\n== 13. Preservación de columnas manuales y depuración de respaldos ==")
from legal.storage import MAX_BACKUPS, read_excel_or_empty

df_extra = read_excel_or_empty(MANDATOS_XLSX)
df_extra["Columna_Manual_Usuario"] = "nota del abogado"
write_excel_safe(df_extra, MANDATOS_XLSX)
df_reread = read_excel_or_empty(MANDATOS_XLSX)
check("Columnas agregadas a mano en Excel se preservan tras reescritura",
      "Columna_Manual_Usuario" in df_reread.columns
      and df_reread["Columna_Manual_Usuario"].iloc[0] == "nota del abogado")
n_backups = len(list(BACKUPS_DIR.glob("*_backup_*")))
check("Respaldos depurados bajo el límite", n_backups <= MAX_BACKUPS)

# ---------------------------------------------------------------------------
print("\n" + "=" * 60)
fallas = [r for r in RESULTADOS if not r[1]]
print(f" RESULTADO FINAL: {len(RESULTADOS) - len(fallas)}/{len(RESULTADOS)} pruebas PASAN")
print("=" * 60)
sys.exit(1 if fallas else 0)
